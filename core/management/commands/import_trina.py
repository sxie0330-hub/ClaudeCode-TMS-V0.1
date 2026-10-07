"""Import the Trina NJ/VA tracking workbook as test data.

    python manage.py import_trina "/path/to/Trina NJ_VA units Tracking Sheet.xlsx" [--dry-run]

Grain: one shipment per container per project/yard. Trina's own BESS boxes
(CYMU…) map 1:1; the KOCU… ocean containers carried BIC units that were later
trucked to different sites one by one, so those split per project. Each gets
  leg 1  OTR  Trina CA → NJ/VA yard   (from the "… CA to NJ/VA" weekly tabs)
  yard stay at NJ/VA yard
  leg 2  OTR  yard → job site         (from the yard-to-site tabs / inventory list)
Re-running is safe: containers already imported are skipped.
"""
import re
from collections import defaultdict
from datetime import date, datetime

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from core.models import CargoLine, Carrier, Customer, Leg, Location, Project, Shipment, Unit, YardStay

YARDS = {
    "NJ": dict(code="NJ Yard", name="NJ Yard", kind=Location.Kind.YARD, address="250 Port St", city="Newark", state="NJ", postal_code="07114"),
    "VA": dict(code="VA Yard", name="VA Yard", kind=Location.Kind.YARD, address="33549 Carver Rd", city="Franklin", state="VA", postal_code="23851"),
}
ORIGIN = dict(code="Trina CA", name="Trina Storage (CA)", kind=Location.Kind.SHIPPER, notes="原表未记录 CA 起运地址，请补充")
CNTR_RE = re.compile(r"^[A-Z]{4}\d{7}$")


def d(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def s(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    v = str(v).strip()
    return "" if v in ("//", "/", "N/A", "NA") else v


def header_map(row):
    return {s(v).lower(): i for i, v in enumerate(row) if s(v)}


def rows_with_headers(ws):
    """Yield dict rows from sheets whose header row repeats in weekly blocks."""
    cols = None
    for row in ws.iter_rows(values_only=True):
        cells = [s(v) for v in row]
        if "container #" in [c.lower() for c in cells] or "container number" in [c.lower() for c in cells]:
            cols = header_map(row)
            continue
        if not cols:
            continue
        rec = {k: row[i] for k, i in cols.items() if i < len(row)}
        cntr = s(rec.get("container #") or rec.get("container number")).upper()
        if CNTR_RE.match(cntr):
            rec["_cntr"] = cntr
            yield rec


def parse_site(text):
    """Keep the sheet's address text verbatim; many are incomplete ("MA 01430")
    and need a human to fix them in the address book."""
    return {"address": text.strip(), "notes": "从表格导入，地址请核对补全"}


class Command(BaseCommand):
    help = "Import the Trina NJ/VA units tracking workbook (test data)."

    def add_arguments(self, parser):
        parser.add_argument("path")
        parser.add_argument("--dry-run", action="store_true")

    def handle(self, path, dry_run=False, **opts):
        try:
            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        except Exception as e:
            raise CommandError(f"无法打开 {path}: {e}")

        c = defaultdict(dict)  # (container, project, yard) -> merged facts
        units = defaultdict(dict)  # same key -> {serial: facts}

        def key_for(cntr, project, yard=None):
            if yard is None:  # tabs that don't say which yard: reuse a known key
                for k in c:
                    if k[0] == cntr and k[1] == project:
                        return k
                yard = "VA"
            return (cntr, project, yard)

        for title, yard in (("NJ Yard Inventory List", "NJ"), ("VA Yard Inventory List", "VA")):
            if title not in wb.sheetnames:
                continue
            for r in rows_with_headers(wb[title]):
                k = key_for(r["_cntr"], s(r.get("project allocation")) or "TBD", yard)
                f = c[k]
                site = s(r.get("site address"))
                f.update(yard=yard, project=k[1], site=site if site.upper() != "TBD" else f.get("site", ""),
                         est_delivery=d(r.get("estimate site delivery date")) or f.get("est_delivery"),
                         actual_delivery=d(r.get("actual delivery date")) or f.get("actual_delivery"),
                         src=f"{title}")
                serial = s(r.get("serial number"))
                if serial:
                    units[k][serial] = dict(product=s(r.get("product")), config=s(r.get("cabinet type")), tier=s(r.get("cell tier")))

        for title in wb.sheetnames:
            m = re.search(r"CA to (NJ|VA)", title)
            if not m:
                continue
            for r in rows_with_headers(wb[title]):
                k = key_for(r["_cntr"], s(r.get("project name")) or "TBD", m.group(1))
                f = c[k]
                f.setdefault("yard", k[2])
                f.setdefault("project", k[1])
                f["in_pu"] = d(r.get("pu date"))
                f["in_del"] = d(r.get("del date"))
                notes = [
                    f"Tarp: {s(r.get('tarp condition'))}" if s(r.get("tarp condition")) else "",
                    f"IIR: {s(r.get('iir result'))}" if s(r.get("iir result")) else "",
                    f"Outbound picture: {s(r.get('outbound picture'))}" if s(r.get("outbound picture")) else "",
                    s(r.get("note")),
                ]
                f["in_notes"] = " · ".join(n for n in notes if n)
                serial = s(r.get("serial no."))
                if serial:
                    units[k].setdefault(serial, dict(config=s(r.get("type")), tier=s(r.get("tier")), product="Elementa"))

        for title in wb.sheetnames:
            if "CA to" in title or "Inventory" in title:
                continue
            ws = wb[title]
            for r in rows_with_headers(ws):
                k = key_for(r["_cntr"], s(r.get("project name")) or "TBD")
                f = c[k]
                f.setdefault("project", k[1])
                f["out_pu"] = d(r.get("pu date")) or f.get("out_pu")
                f["out_sched"] = d(r.get("schedule del date")) or f.get("out_sched")
                f["actual_delivery"] = d(r.get("actual delivery date")) or d(r.get("del date")) or f.get("actual_delivery")
                f["out_carrier"] = s(r.get("carrier")) or f.get("out_carrier", "")
                f.setdefault("yard", k[2])
                serial = s(r.get("serial no."))
                if serial:
                    units[k].setdefault(serial, dict(config=s(r.get("type")), tier=s(r.get("tier")), product=""))

        self.stdout.write(f"读到 {len(c)} 个柜子，{sum(len(v) for v in units.values())} 个序列号")
        if dry_run:
            for k, f in list(c.items())[:10]:
                self.stdout.write(f"  {k}: {f} units={list(units[k])}")
            return

        created = skipped = 0
        with transaction.atomic():
            customer, _ = Customer.objects.get_or_create(name="Trina", defaults={"code": "TRINA"})
            yards = {key: Location.objects.get_or_create(code=v["code"], defaults=v)[0] for key, v in YARDS.items()}
            origin, _ = Location.objects.get_or_create(code=ORIGIN["code"], defaults=ORIGIN)
            projects, sites, carriers = {}, {}, {}
            for k, f in sorted(c.items(), key=lambda kv: (kv[1].get("in_pu") or date.max, kv[0])):
                cntr, pname, yard_key = k
                source = f"Trina NJ_VA units Tracking Sheet.xlsx · {cntr} · {pname} · {yard_key} Yard"
                if Shipment.all_objects.filter(source_ref=source).exists():
                    skipped += 1
                    continue
                site = None
                if f.get("site"):
                    if f["site"] not in sites:
                        sites[f["site"]] = Location.objects.get_or_create(
                            code=f"Trina-{pname}"[:60], defaults={"name": f"{pname} Job Site", "kind": Location.Kind.SITE, **parse_site(f["site"])}
                        )[0]
                    site = sites[f["site"]]
                if pname not in projects:
                    projects[pname] = Project.objects.filter(customer=customer, name=pname).first() or Project.objects.create(
                        customer=customer, name=pname, default_shipper=origin, default_consignee=site,
                        cargo_description="Energy Storage System (Trina Elementa)",
                        emergency_contact="CHEMTREC 703-527-3887",
                    )
                project = projects[pname]
                yard = yards[f.get("yard", "NJ")]
                sh = Shipment.objects.create(
                    customer=customer, project=project, service=Shipment.Service.OTR, container_number=cntr,
                    equipment=Shipment.Equipment.C20SOC if cntr.startswith("CYMU") else "", shipper=origin, consignee=site or project.default_consignee,
                    source_ref=source,
                )
                for serial, u in units[k].items():
                    Unit.objects.create(shipment=sh, serial_number=serial, product=u.get("product") or "Elementa",
                                        config=u.get("config", ""), tier=u.get("tier", ""))
                CargoLine.objects.create(shipment=sh, handling_qty=1, handling_type=CargoLine.HandlingType.UNIT,
                                         description="Energy Storage System (Trina Elementa)")
                if f.get("in_pu") or f.get("in_del"):
                    Leg.objects.create(shipment=sh, sequence=1, kind=Leg.Kind.OTR, origin=origin, destination=yard,
                                       actual_pickup=f.get("in_pu"), actual_delivery=f.get("in_del"), notes=f.get("in_notes", ""))
                out_pu = f.get("out_pu")
                delivered = f.get("actual_delivery")
                if out_pu is None and delivered:
                    out_pu = delivered  # original sheet only kept the delivery date
                if f.get("in_del") or out_pu:
                    YardStay.objects.create(shipment=sh, yard=yard, in_date=f.get("in_del"), out_date=out_pu)
                if site or f.get("est_delivery") or f.get("out_sched") or delivered:
                    carrier = None
                    if f.get("out_carrier"):
                        name = f["out_carrier"]
                        carriers[name] = carriers.get(name) or Carrier.objects.get_or_create(name=name)[0]
                        carrier = carriers[name]
                    Leg.objects.create(shipment=sh, sequence=2, kind=Leg.Kind.OTR, origin=yard, destination=site, carrier=carrier,
                                       scheduled_delivery=f.get("out_sched") or f.get("est_delivery"),
                                       actual_pickup=out_pu if delivered or f.get("out_pu") else None, actual_delivery=delivered)
                sh.refresh_status()
                created += 1
        self.stdout.write(self.style.SUCCESS(f"导入完成：新建 {created} 票，跳过已存在 {skipped} 票"))
